#!/usr/bin/env python3
"""Build the survey supporting workbook in public/survey/."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import csv

ROOT = Path(__file__).resolve().parents[1]
SURVEY = ROOT / "public" / "survey"
OUT = SURVEY / "Wishlist-survey-workbook.xlsx"

INK = "1C1917"
PAPER = "FFFCF8"
SUNKEN = "EBE4DA"
ACCENT = "C21845"

thin = Border(
    left=Side(style="thin", color="C4BDB3"),
    right=Side(style="thin", color="C4BDB3"),
    top=Side(style="thin", color="C4BDB3"),
    bottom=Side(style="thin", color="C4BDB3"),
)
head_fill = PatternFill("solid", fgColor="1C1917")
head_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
title_font = Font(name="Calibri", bold=True, size=16, color=INK)
body = Font(name="Calibri", size=11, color=INK)
wrap = Alignment(wrap_text=True, vertical="top")


def style_header(ws, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(1, c)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = thin


def autosize(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def read_csv(name):
    with (SURVEY / name).open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def main():
    inst = read_csv("instrument.csv")
    code = read_csv("codebook.csv")
    responses = read_csv("responses.csv")
    ids = [r["id"] for r in inst]
    n_resp = len(responses)
    last = n_resp + 1
    spare = last + 10

    wb = Workbook()

    # --- Read_me ---
    ws = wb.active
    ws.title = "Read_me"
    ws.sheet_properties.tabColor = ACCENT
    lines = [
        ("Verdict — wishlist survey workbook", title_font),
        ("", body),
        ("What this is", Font(name="Calibri", bold=True, size=12, color=INK)),
        (
            "Supporting artefact for Part 3. Interviews stay the qualitative lock. This workbook is the survey database. The live fillable form is /survey/form/ on the case site. Responses tab is the panel for this case.",
            body,
        ),
        ("", body),
        ("What is in Responses", Font(name="Calibri", bold=True, size=11, color=ACCENT)),
        (
            "Constructed case sample (same method as the six interview write-ups). Cities: Bengaluru, Hyderabad, Delhi, Gurugram, Noida, Coimbatore, plus out-of-target noise (Mumbai, Pune, non-metro, other ages/genders). in_target and q3_code fill by formula.",
            body,
        ),
        ("", body),
        ("Host", Font(name="Calibri", bold=True, size=12, color=INK)),
        ("Live form: case site /survey/form/. Database = this xlsx.", body),
        ("", body),
        ("How the pieces fit", Font(name="Calibri", bold=True, size=12, color=INK)),
        ("1. Instrument tab = every question. The live form uses the same IDs (S1. … Q3. …).", body),
        ("2. Responses tab = the case panel evaluators open.", body),
        ("3. Keep formula columns in_target / q3_code / off_app / legal_lever. Do not overwrite them.", body),
        ("", body),
        ("in_target rule", Font(name="Calibri", bold=True, size=12, color=INK)),
        ("Woman AND age 24–32 AND city in Bengaluru / Hyderabad / Delhi / Gurugram / Noida / Coimbatore AND wishlist 15+ AND delayed a save in 60 days AND opens Myntra weekly or more.", body),
        ("", body),
        ("If sale-wait is the most common Q3 in target", Font(name="Calibri", bold=True, size=12, color=INK)),
        ("Rank it. Do not ship a coupon. Check Q9 for ‘wait for a sale so a failed return hurts less’ — that is return-fear wearing a discount.", body),
        ("", body),
        ("Tabs", Font(name="Calibri", bold=True, size=12, color=INK)),
        ("Instrument — every question, type, options, brief mapping.", body),
        ("Codebook — how to recode Q3 / Q5 / Q8 / Q12.", body),
        ("Responses — the database. Row 1 is field IDs. Sample rows filled; spare rows below for live paste.", body),
        ("Analysis — counts that update from in_target / q3_code.", body),
    ]
    ws.column_dimensions["A"].width = 110
    for i, (text, font) in enumerate(lines, 1):
        cell = ws.cell(i, 1, text)
        cell.font = font
        cell.alignment = wrap
        ws.row_dimensions[i].height = 18 if len(text) < 80 else 48
    ws.row_dimensions[4].height = 72
    ws.row_dimensions[7].height = 48

    # --- Instrument ---
    wi = wb.create_sheet("Instrument")
    headers = list(inst[0].keys())
    wi.append(headers)
    for row in inst:
        wi.append([row[h] for h in headers])
    style_header(wi, len(headers))
    autosize(wi, [10, 28, 12, 10, 70, 70, 36, 22])
    for r in wi.iter_rows(min_row=2, max_row=wi.max_row, max_col=len(headers)):
        for c in r:
            c.alignment = wrap
            c.font = body
            c.border = thin
        wi.row_dimensions[r[0].row].height = 36
    wi.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{wi.max_row}"
    wi.freeze_panes = "A2"

    # --- Codebook ---
    wc = wb.create_sheet("Codebook")
    ch = list(code[0].keys())
    wc.append(ch)
    for row in code:
        wc.append([row[h] for h in ch])
    style_header(wc, len(ch))
    autosize(wc, [12, 55, 70, 55])
    for r in wc.iter_rows(min_row=2, max_row=wc.max_row, max_col=len(ch)):
        for c in r:
            c.alignment = wrap
            c.font = body
            c.border = thin
        wc.row_dimensions[r[0].row].height = 40
    wc.freeze_panes = "A2"

    # --- Responses ---
    wr = wb.create_sheet("Responses")
    extra = ["in_target", "q3_code", "off_app", "legal_lever"]
    resp_headers = ["Timestamp", *ids, *extra]
    wr.append(resp_headers)
    style_header(wr, len(resp_headers))
    # Column index map
    col = {h: i + 1 for i, h in enumerate(resp_headers)}

    def L(name, row):
        return f"{get_column_letter(col[name])}{row}"

    for r, rec in enumerate(responses, 2):
        wr.cell(r, 1, rec.get("Timestamp", ""))
        for fid in ids:
            wr.cell(r, col[fid], rec.get(fid, ""))
        wr.row_dimensions[r].height = 18

    def fill_formulas(r: int):
        cities = (
            f'OR({L("S4", r)}="Bengaluru",{L("S4", r)}="Hyderabad",{L("S4", r)}="Delhi",'
            f'{L("S4", r)}="Gurugram",{L("S4", r)}="Noida",{L("S4", r)}="Coimbatore")'
        )
        wr.cell(
            r,
            col["in_target"],
            f'=IF({L("S1", r)}="","",IF(AND({L("S2", r)}="Woman",{L("S3", r)}="24–32",{cities},OR({L("S5", r)}="15–40",{L("S5", r)}="41–80",{L("S5", r)}="81+"),{L("S6", r)}="Yes",OR({L("S7", r)}="Most days",{L("S7", r)}="About weekly")),"YES","NO"))',
        )
        wr.cell(
            r,
            col["q3_code"],
            f'=IF({L("Q3", r)}="","",IF(ISNUMBER(SEARCH("fit",{L("Q3", r)})),"fit",IF(ISNUMBER(SEARCH("return",{L("Q3", r)}))+ISNUMBER(SEARCH("seal",{L("Q3", r)})),"return",IF(ISNUMBER(SEARCH("chart",{L("Q3", r)})),"chart",IF(ISNUMBER(SEARCH("choose",{L("Q3", r)})),"compare",IF(ISNUMBER(SEARCH("sale",{L("Q3", r)})),"sale",IF(ISNUMBER(SEARCH("never",{L("Q3", r)})),"bookmark",IF(ISNUMBER(SEARCH("mess",{L("Q3", r)})),"clutter",IF(ISNUMBER(SEARCH("wear it",{L("Q3", r)})),"occasion",IF(ISNUMBER(SEARCH("quality",{L("Q3", r)})),"quality","other"))))))))))',
        )
        wr.cell(
            r,
            col["off_app"],
            f'=IF({L("Q8", r)}="","",IF(OR(ISNUMBER(SEARCH("WhatsApp",{L("Q8", r)})),ISNUMBER(SEARCH("YouTube",{L("Q8", r)})),ISNUMBER(SEARCH("Instagram",{L("Q8", r)})),ISNUMBER(SEARCH("Google",{L("Q8", r)})),ISNUMBER(SEARCH("store",{L("Q8", r)})),ISNUMBER(SEARCH("two sizes",{L("Q8", r)}))),"off_app","in_app"))',
        )
        wr.cell(
            r,
            col["legal_lever"],
            f'=IF({L("Q5", r)}="","",IF(ISNUMBER(SEARCH("coupon",{L("Q5", r)}))+ISNUMBER(SEARCH("sale",{L("Q5", r)}))+ISNUMBER(SEARCH("kept-vs-returned",{L("Q5", r)}))+ISNUMBER(SEARCH("sentence",{L("Q5", r)}))+ISNUMBER(SEARCH("stays with",{L("Q5", r)}))+ISNUMBER(SEARCH("side-by-side",{L("Q5", r)})),IF(ISNUMBER(SEARCH("kept-vs-returned",{L("Q5", r)}))+ISNUMBER(SEARCH("sentence",{L("Q5", r)}))+ISNUMBER(SEARCH("stays with",{L("Q5", r)}))+ISNUMBER(SEARCH("side-by-side",{L("Q5", r)})),"legal_present",IF(ISNUMBER(SEARCH("coupon",{L("Q5", r)}))+ISNUMBER(SEARCH("sale",{L("Q5", r)})),"coupon_only","other")),"other"))',
        )

    for r in range(2, spare + 1):
        fill_formulas(r)
        wr.row_dimensions[r].height = 18

    for c in range(1, len(resp_headers) + 1):
        wr.cell(1, c).border = thin
        wr.column_dimensions[get_column_letter(c)].width = 18
    wr.column_dimensions["A"].width = 20
    wr.freeze_panes = "B2"
    wr.auto_filter.ref = f"A1:{get_column_letter(len(resp_headers))}{spare}"
    wr.sheet_properties.tabColor = "3D5AFE"

    note = wr.cell(
        spare + 1,
        1,
        "Sample rows are filled. Paste live Form answers over Timestamp…Q13. Do not delete formula columns in_target / q3_code / off_app / legal_lever.",
    )
    note.font = Font(name="Calibri", italic=True, size=10, color="78716C")
    wr.merge_cells(start_row=spare + 1, start_column=1, end_row=spare + 1, end_column=8)

    # --- Analysis ---
    wa = wb.create_sheet("Analysis")
    wa["A1"] = "Cuts that update from Responses (target = in_target YES)"
    wa["A1"].font = title_font
    wa.merge_cells("A1:C1")
    headers_a = ["Cut", "How to read it", "Count / share"]
    for i, h in enumerate(headers_a, 1):
        cell = wa.cell(3, i, h)
        cell.fill = head_fill
        cell.font = head_font
        cell.border = thin

    it = get_column_letter(col["in_target"])
    q3c = get_column_letter(col["q3_code"])
    q2 = get_column_letter(col["Q2"])
    q10 = get_column_letter(col["Q10a"])
    q12 = get_column_letter(col["Q12"])

    cuts = [
        ("N with timestamp", "Started the form", f'=COUNTA(Responses!A2:A{spare})'),
        ("N in_target", "Woman 24–32 metro 15+ weekly delayed save", f'=COUNTIF(Responses!{it}2:{it}{spare},"YES")'),
        ("Target · still intend Yes", "North-star population inside target", f'=COUNTIFS(Responses!{it}2:{it}{spare},"YES",Responses!{q2}2:{q2}{spare},"Yes")'),
        ("Target · Q3 = fit", "Legal primary freeze", f'=COUNTIFS(Responses!{it}2:{it}{spare},"YES",Responses!{q3c}2:{q3c}{spare},"fit")'),
        ("Target · Q3 = return", "Co-primary", f'=COUNTIFS(Responses!{it}2:{it}{spare},"YES",Responses!{q3c}2:{q3c}{spare},"return")'),
        ("Target · Q3 = sale", "Rank. Do not ship.", f'=COUNTIFS(Responses!{it}2:{it}{spare},"YES",Responses!{q3c}2:{q3c}{spare},"sale")'),
        ("Target · Q3 = bookmark", "Still exploring. Do not convert.", f'=COUNTIFS(Responses!{it}2:{it}{spare},"YES",Responses!{q3c}2:{q3c}{spare},"bookmark")'),
        ("Target · PDP story vanished (Q10a Yes)", "Mechanism for Verdict on the card", f'=COUNTIFS(Responses!{it}2:{it}{spare},"YES",Responses!{q10}2:{q10}{spare},"Yes")'),
        ("Target · wishlist is moodboard (Q12)", "Over-nudge guardrail", f'=COUNTIFS(Responses!{it}2:{it}{spare},"YES",Responses!{q12}2:{q12}{spare},"A moodboard / Pinterest")'),
    ]

    for i, (cut, how, formula) in enumerate(cuts, 4):
        wa.cell(i, 1, cut).font = body
        wa.cell(i, 2, how).alignment = wrap
        wa.cell(i, 3, formula)
        for c in range(1, 4):
            wa.cell(i, c).border = thin
            wa.cell(i, c).font = body

    # % fit among target = fit count / in_target  → rows 5 and 7? 
    # Row 4 = N timestamp
    # Row 5 = N in_target
    # Row 6 = intend
    # Row 7 = fit
    wa.cell(13, 1, "% of target whose Q3 is fit")
    wa.cell(13, 2, "If sale > fit, still do not coupon. Check Q9 for risk-offset.")
    wa.cell(13, 3, '=IF(C5=0,"",C7/C5)')
    wa.cell(14, 1, "% of target whose Q3 is sale")
    wa.cell(14, 2, "DISQUALIFIED lever frequency")
    wa.cell(14, 3, '=IF(C5=0,"",C9/C5)')
    for r in (13, 14):
        wa.cell(r, 3).number_format = "0.0%"
        for c in range(1, 4):
            wa.cell(r, c).border = thin
            wa.cell(r, c).font = body
            wa.cell(r, c).alignment = wrap

    wa.column_dimensions["A"].width = 42
    wa.column_dimensions["B"].width = 48
    wa.column_dimensions["C"].width = 22
    wa.row_dimensions[1].height = 24
    wa["A16"] = "Share this xlsx (or the live Google Sheet view link) as the supporting URL. Form URL is a slot on Docs until you paste it."
    wa["A16"].font = Font(name="Calibri", italic=True, size=10, color="78716C")
    wa.merge_cells("A16:C16")

    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
